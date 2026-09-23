"""Local, session-independent bibliographic ingestion, cleaning and PRISMA exports."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import BinaryIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from citation_readers import CitationParseError, read_citations
from utils import clean_link_series, clean_series, decode_text, normalize_column, normalize_text, unique_headers


FIELDS = [
    "Titulo", "Autores", "Año", "DOI", "Resumen", "Palabras_clave", "Revista",
    "URL", "Tipo_documento", "Afiliacion",
]
EXTRA_FIELDS = ["Serie_libro", "Volumen", "Numero", "Paginas", "ISSN", "Fecha_publicacion"]
SUPPORTED_FIELDS = FIELDS + EXTRA_FIELDS
SCREENING_FIELDS = [
    "Decision_Titulo_Resumen", "Motivo_Exclusion", "Texto_Completo", "Decision_Final",
]
FIELD_LABELS = {
    "Titulo": "Título", "Autores": "Autores", "Año": "Año", "DOI": "DOI",
    "Resumen": "Resumen", "Palabras_clave": "Palabras clave", "Revista": "Revista o fuente",
    "URL": "URL", "Tipo_documento": "Tipo de documento", "Afiliacion": "Afiliación",
    "Serie_libro": "Serie de libros", "Volumen": "Volumen", "Numero": "Número",
    "Paginas": "Páginas o número de artículo", "ISSN": "ISSN", "Fecha_publicacion": "Fecha de publicación",
}
ALIASES = {
    "Titulo": ("title", "document title", "article title", "paper title", "item title", "titulo", "ti", "t1"),
    "Autores": ("authors", "author", "author(s)", "creators", "creator", "author full names", "autores", "autor", "au", "af", "a1"),
    "Año": ("year", "publication year", "publication date", "date", "cover date", "published", "online publication date", "print publication date", "año", "fecha de publicacion", "py", "y1", "da"),
    "DOI": ("doi", "digital object identifier", "document doi", "article doi", "item doi", "di", "do"),
    "Resumen": ("abstract", "summary", "description", "abstract note", "resumen", "ab", "n2"),
    "Palabras_clave": ("keywords", "author keywords", "index keywords", "keyword", "key words", "keywords plus", "palabras clave", "de", "kw"),
    "Revista": ("journal", "journal title", "source title", "publication title", "source", "publication", "container title", "journal book", "revista", "fuente", "so", "jo", "jf", "t2"),
    "URL": ("url", "link", "document url", "article url", "web link", "url link", "enlace", "ur"),
    "Tipo_documento": ("document type", "publication type", "type", "article type", "content type", "tipo documento", "tipo de documento", "dt", "ty"),
    "Afiliacion": ("affiliation", "affiliations", "author affiliations", "authors affiliations", "institution", "organization", "addresses", "afiliacion", "afiliaciones", "institucion", "c1", "ad"),
    "Serie_libro": ("book series title", "series", "series title", "serie libro", "serie de libros", "t3"),
    "Volumen": ("journal volume", "volume", "volumen", "vl"),
    "Numero": ("journal issue", "issue", "number", "numero", "is"),
    "Paginas": ("pages", "page range", "article number", "paginas"),
    "ISSN": ("issn", "print issn", "electronic issn", "eissn", "sn"),
    "Fecha_publicacion": ("publication date", "date", "cover date", "online publication date", "print publication date", "fecha publicacion", "fecha de publicacion", "da"),
}
MAX_FILE_BYTES = 50 * 1024 * 1024
DOI_PREFIX = r"^(?:(?:https?://)?(?:dx\.)?doi\.org/|doi\s*:\s*)+"
YEAR_PATTERN = r"(?<!\d)((?:1[5-9]|20|21)\d{2})(?!\d)"


class FileProcessingError(ValueError):
    """An expected input problem that can be shown directly to the user."""


class ColumnDetectionError(FileProcessingError):
    """No supported bibliographic column was found; manual mapping is possible."""


@dataclass
class ProcessingResult:
    """Normalized records, auditable duplicates and an independent screening base."""

    records: pd.DataFrame
    clean: pd.DataFrame
    duplicates: pd.DataFrame
    metrics: dict[str, int]
    mapping: dict[str, list[str]]
    warnings: list[str]


def _read_csv(data: bytes) -> tuple[pd.DataFrame, dict]:
    text, encoding = decode_text(data)
    if not text or not text.strip():
        raise FileProcessingError("El archivo está vacío. Sube una exportación con encabezados y registros.")
    if "\x00" in text or any(ord(char) < 9 for char in text[:8192]):
        raise FileProcessingError("El archivo no pudo ser interpretado como CSV. Guárdalo como CSV UTF-8 e inténtalo de nuevo.")
    text = text.lstrip("\ufeff\r\n")
    first_line = text.splitlines()[0].strip()
    if first_line.casefold() in ("sep=,", "sep=;", "sep="):
        delimiter = first_line[-1] if first_line != "sep=" else "\t"
        text = text.split("\n", 1)[1] if "\n" in text else ""
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t").delimiter
        except csv.Error:
            # A one-column export is valid; a malformed multi-column export is
            # still checked below instead of silently dropping its records.
            delimiter = max(",;\t", key=lambda sep: len(next(csv.reader([first_line], delimiter=sep))))
    rows = []
    try:
        for row in csv.reader(StringIO(text), delimiter=delimiter, strict=True):
            if row and any(cell.strip() for cell in row):
                if rows and len(row) != len(rows[0]):
                    raise FileProcessingError(
                        "El CSV contiene filas con un número distinto de columnas. "
                        "Revisa el separador y las comillas, o vuelve a exportarlo como XLSX."
                    )
                rows.append(row)
    except csv.Error as exc:
        raise FileProcessingError("El CSV tiene comillas o delimitadores inconsistentes. Revisa el archivo o expórtalo como XLSX.") from exc
    if len(rows) < 2:
        raise FileProcessingError("El archivo no contiene registros. Incluye encabezados y al menos una fila de datos.")
    frame = pd.DataFrame(rows[1:], columns=unique_headers(rows[0]))
    return frame, {"encoding": encoding, "separator": delimiter}


def _read_excel(data: bytes) -> tuple[pd.DataFrame, dict]:
    try:
        with pd.ExcelFile(BytesIO(data), engine="openpyxl") as workbook:
            for sheet in workbook.sheet_names:
                raw = pd.read_excel(workbook, sheet_name=sheet, header=None, dtype=str, keep_default_na=False)
                if raw.empty:
                    continue
                raw = raw.loc[~raw.apply(clean_series).eq("").all(axis=1)]
                if len(raw) < 2:
                    continue
                frame = raw.iloc[1:].copy()
                frame.columns = unique_headers(raw.iloc[0].tolist())
                return frame, {"sheet": sheet, "sheets": workbook.sheet_names}
    except (BadZipFile, InvalidFileException, ParseError, OSError, ValueError, KeyError) as exc:
        raise FileProcessingError("El archivo no pudo ser interpretado como Excel. Verifica que no esté dañado y que sea un archivo XLSX.") from exc
    raise FileProcessingError("El Excel no contiene datos. Incluye encabezados y al menos un registro en una hoja.")


def read_uploaded_file(uploaded_file: bytes | BinaryIO, filename: str | None = None) -> pd.DataFrame:
    """Read CSV, XLSX, RIS, BibTeX or structured citation text in memory.

    Input metadata is attached to DataFrame.attrs. No uploaded file is persisted.
    CSV rows with inconsistent widths are rejected, never silently skipped.
    """
    name = filename or getattr(uploaded_file, "name", "")
    extension = Path(name).suffix.lower()
    if extension not in {".csv", ".xlsx", ".ris", ".bib", ".bibtex", ".txt"}:
        raise FileProcessingError("Formato no compatible. Sube un archivo CSV, XLSX, RIS, BibTeX o TXT.")
    if isinstance(uploaded_file, bytes):
        data = uploaded_file
    elif hasattr(uploaded_file, "getvalue"):
        data = uploaded_file.getvalue()
    else:
        uploaded_file.seek(0)
        data = uploaded_file.read()
    if not data:
        raise FileProcessingError("El archivo está vacío. Selecciona una exportación con registros.")
    if len(data) > MAX_FILE_BYTES:
        raise FileProcessingError("El archivo supera los 50 MB. Divide la exportación en archivos más pequeños.")
    if extension == ".csv":
        frame, metadata = _read_csv(data)
    elif extension == ".xlsx":
        frame, metadata = _read_excel(data)
    else:
        text, encoding = decode_text(data)
        expected = {".ris": "RIS", ".bib": "BibTeX", ".bibtex": "BibTeX"}.get(extension)
        try:
            frame, metadata = read_citations(text, expected)
        except CitationParseError as exc:
            raise FileProcessingError(str(exc)) from exc
        metadata["encoding"] = encoding
    cleaned = frame.apply(clean_series)
    empty = cleaned.eq("").all(axis=1)
    frame = frame.loc[~empty].reset_index(drop=True)
    if frame.empty:
        raise FileProcessingError("No se encontraron registros: todas las filas están vacías.")
    frame.attrs.update(metadata, filename=name, file_type=extension[1:].upper(), empty_rows_removed=int(empty.sum()))
    return frame


def detect_columns(dataframe: pd.DataFrame) -> dict[str, list[str]]:
    """Match aliases without relying on column order; retain repeated source fields.

    More specific, preferred aliases win when several populated columns match.
    Secondary matches fill missing cells. Keyword columns are combined separately.
    """
    normalized = {
        column: normalize_column(re.sub(r" \[columna \d+\]$", "", str(column)))
        for column in dataframe.columns
    }
    mapping = {}
    for target, aliases in ALIASES.items():
        matches = []
        for alias in aliases:
            for column, name in normalized.items():
                if name == normalize_column(alias) and column not in matches:
                    matches.append(column)
        mapping[target] = matches
    return mapping


def normalize_doi(value: object) -> str:
    """Remove common resolver URLs/DOI prefixes and normalize case; never invent a DOI."""
    return normalize_doi_series(pd.Series([value])).iloc[0]


def normalize_doi_series(series: pd.Series) -> pd.Series:
    """Use one DOI normalization path for every import and duplicate comparison."""
    values = clean_series(series).str.replace(r"(?i)^(?:doi\s*:\s*)+", "", regex=True)
    return clean_link_series(values).str.lower().str.replace(DOI_PREFIX, "", regex=True).str.strip()


def extract_year(value: object) -> str:
    """Extract an unambiguous four-digit year from a year or publication date."""
    matches = set(re.findall(YEAR_PATTERN, normalize_text(value)))
    return matches.pop() if len(matches) == 1 else ""


def normalize_titles(series: pd.Series) -> pd.Series:
    """Produce punctuation-free comparison keys while keeping displayed titles intact."""
    return (
        clean_series(series).str.normalize("NFKC").str.casefold()
        .str.replace(r"[\W_]+", " ", regex=True).str.replace(r"\s+", " ", regex=True).str.strip()
    )


def find_duplicates(dataframe: pd.DataFrame) -> pd.DataFrame:
    """Flag later occurrences by nonempty DOI and/or title, in original file order."""
    result = dataframe.copy()
    doi = normalize_doi_series(result["DOI"])
    titles = normalize_titles(result["Titulo"])
    duplicate_doi = doi.ne("") & doi.duplicated(keep="first")
    duplicate_title = titles.ne("") & titles.duplicated(keep="first")
    result["Duplicado"] = duplicate_doi | duplicate_title
    result["Motivo_duplicado"] = ""
    result.loc[duplicate_doi, "Motivo_duplicado"] = "DOI duplicado"
    result.loc[duplicate_title, "Motivo_duplicado"] = "Título duplicado"
    result.loc[duplicate_doi & duplicate_title, "Motivo_duplicado"] = "DOI y título duplicados"
    return result


def create_prisma_dataframe(records: pd.DataFrame) -> pd.DataFrame:
    """Create the unique-record screening base, keeping the first occurrence."""
    columns = [field for field in SUPPORTED_FIELDS if field in records.columns]
    clean = records.loc[~records["Duplicado"], columns].reset_index(drop=True).copy()
    clean.insert(0, "ID", range(1, len(clean) + 1))
    for field in SCREENING_FIELDS:
        clean[field] = ""
    return clean


def calculate_prisma_metrics(records: pd.DataFrame) -> dict[str, int]:
    """Count records after removing empty rows; coverage metrics include duplicates."""
    return {
        "Registros identificados": len(records),
        "Duplicados encontrados": int(records["Duplicado"].sum()),
        "Registros únicos": int((~records["Duplicado"]).sum()),
        "Registros con DOI": int(records["DOI"].ne("").sum()),
        "Registros con abstract": int(records["Resumen"].ne("").sum()),
    }


def process_dataframe(dataframe: pd.DataFrame, overrides: dict[str, str | None] | None = None) -> ProcessingResult:
    """Normalize detected fields and build a complete, auditable PRISMA result."""
    source = dataframe.copy()
    source.columns = unique_headers(list(source.columns))
    source = source.apply(clean_series)
    source = source.loc[~source.eq("").all(axis=1)].reset_index(drop=True)
    if source.empty:
        raise FileProcessingError("No se encontraron registros para procesar.")
    mapping = detect_columns(source)
    for field, column in (overrides or {}).items():
        if field not in SUPPORTED_FIELDS or (column is not None and column not in source.columns):
            raise FileProcessingError("La selección de columnas no corresponde al archivo cargado.")
        mapping[field] = [column] if column is not None else []
    if not any(mapping.values()):
        raise ColumnDetectionError("No se reconocieron columnas bibliográficas. Asigna las columnas en la sección inferior para continuar.")
    columns = FIELDS + [field for field in EXTRA_FIELDS if mapping[field]]
    records = pd.DataFrame("", index=source.index, columns=columns)
    for field, columns in mapping.items():
        for column in columns:
            values = source[column]
            if field == "Palabras_clave":
                add = values.ne("") & records[field].ne("") & records[field].ne(values)
                records.loc[add, field] = records.loc[add, field] + "; " + values.loc[add]
            records[field] = records[field].mask(records[field].eq(""), values)
    records["DOI"] = normalize_doi_series(records["DOI"])
    records["URL"] = clean_link_series(records["URL"])
    recovered = pd.Series(False, index=records.index)
    if not (overrides is not None and "DOI" in overrides and overrides["DOI"] is None):
        # Only an explicit DOI resolver URL can fill a missing DOI. Never scan
        # titles or abstracts, which may cite identifiers belonging to other work.
        recovered = records["DOI"].eq("") & records["URL"].str.fullmatch(
            r"(?i)(?:https?://)?(?:dx\.)?doi\.org/10\.\d{4,9}/\S+", na=False,
        )
        records.loc[recovered, "DOI"] = normalize_doi_series(records.loc[recovered, "URL"])
    years = records["Año"].str.findall(YEAR_PATTERN)
    first_year = years.str[0].fillna("")
    # Remove repeated occurrences of the same year, but do not guess across ranges.
    ambiguous = years.str.len().gt(1)
    if ambiguous.any():
        first_year.loc[ambiguous] = records.loc[ambiguous, "Año"].map(extract_year)
    records["Año"] = first_year
    records = find_duplicates(records)
    warnings = []
    if recovered.any():
        warnings.append(f"Se recuperó el DOI desde enlaces doi.org del campo URL en {int(recovered.sum())} registros.")
    if not mapping["DOI"] and records["DOI"].eq("").all():
        warnings.append("No se detectó una columna de DOI. El procesamiento continuará sin DOI.")
    if not mapping["Resumen"]:
        warnings.append("No se encontró una columna de resumen. Los registros podrán procesarse igualmente.")
    if not mapping["Titulo"]:
        warnings.append("No se detectó una columna de título. Puedes asignarla manualmente; los duplicados se identificarán solo por DOI.")
    unmatchable = int((records["Titulo"].eq("") & records["DOI"].eq("")).sum())
    if unmatchable:
        warnings.append(f"{unmatchable} registros no tienen título ni DOI: se conservaron y requieren revisión manual.")
    return ProcessingResult(
        records=records, clean=create_prisma_dataframe(records),
        duplicates=records.loc[records["Duplicado"]].reset_index(drop=True),
        metrics=calculate_prisma_metrics(records), mapping=mapping, warnings=warnings,
    )


def filter_records(
    records: pd.DataFrame, query: str = "", year_range: tuple[int, int] | None = None,
    doi: str = "Todos", abstract: str = "Todos", include_unknown_year: bool = True,
) -> pd.DataFrame:
    """Filter a view without mutating the full dataset or interpreting search regex."""
    mask = pd.Series(True, index=records.index)
    if query.strip():
        matches = pd.Series(False, index=records.index)
        for field in ("Titulo", "Resumen", "Palabras_clave"):
            matches |= records[field].str.contains(query.strip(), case=False, regex=False, na=False)
        mask &= matches
    if year_range is not None:
        years = pd.to_numeric(records["Año"], errors="coerce")
        mask &= years.between(*year_range) | (years.isna() & include_unknown_year)
    if doi != "Todos":
        mask &= records["DOI"].ne("") if doi == "Con DOI" else records["DOI"].eq("")
    if abstract != "Todos":
        mask &= records["Resumen"].ne("") if abstract == "Con abstract" else records["Resumen"].eq("")
    return records.loc[mask].copy()


def export_csv(dataframe: pd.DataFrame) -> bytes:
    """Create a BOM-prefixed CSV, escaping spreadsheet formula-like text cells."""
    safe = dataframe.copy()
    for column in safe.select_dtypes(include=["object", "string"]).columns:
        values = safe[column].fillna("").astype(str)
        formula = values.str.match(r"^\s*[=+@-]", na=False)
        safe[column] = values.mask(formula, "'" + values)
    return safe.to_csv(index=False).encode("utf-8-sig")


def export_excel(dataframe: pd.DataFrame, metrics: dict[str, int] | None = None, sheet_name: str = "Papers") -> bytes:
    """Export literal cells and a PRISMA summary into an in-memory Excel workbook."""
    if len(dataframe) > 1_048_575:
        raise FileProcessingError("La base supera el límite de filas de Excel. Descarga el CSV.")
    for column in dataframe.select_dtypes(include=["object", "string"]).columns:
        if dataframe[column].fillna("").astype(str).str.len().gt(32767).any():
            raise FileProcessingError("Una celda supera los 32.767 caracteres permitidos por Excel. Descarga el CSV para conservar el contenido completo.")
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Excel disallows certain control characters. They are removed only in
        # the workbook; the normalized in-memory records remain unchanged.
        excel_data = dataframe.replace(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", regex=True)
        excel_data.to_excel(writer, sheet_name=sheet_name, index=False)
        if metrics is not None:
            pd.DataFrame(metrics.items(), columns=["Métrica", "Cantidad"]).to_excel(writer, sheet_name="Resumen_PRISMA", index=False)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.data_type = "s"
            for cell in worksheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="087E8B")
                cell.alignment = Alignment(vertical="center")
                worksheet.column_dimensions[cell.column_letter].width = 48 if cell.value in {"Titulo", "Resumen", "Métrica"} else 24
            worksheet.row_dimensions[1].height = 25
    return buffer.getvalue()
