"""Behavioral checks for diverse exports, deduplication and download integrity."""

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from processor import (
    FIELDS, SCREENING_FIELDS, ColumnDetectionError, FileProcessingError,
    detect_columns, export_csv, export_excel, extract_year, filter_records,
    normalize_doi, normalize_text, process_dataframe, read_uploaded_file,
)
from utils import clean_series, unique_headers


@pytest.fixture
def bibliography():
    return pd.DataFrame({
        "Document Title": ["Deep Learning for Debris Flow", "deep learning for debris-flow", "Another title", "Flood model", "Untimed paper"],
        "Authors": ["  Ana  Pérez\n", "Ana Pérez", None, "Lee", "Chen"],
        "Publication Year": ["2025-03-20", "April 2025", "2024", "20/04/2023", "unknown"],
        "DOI": ["https://doi.org/10.1000/ABC", "10.1000/abc", "DOI: 10.1000/ABC", "", "10.1000/other"],
        "Abstract": ["Neural networks", "Same paper", "Other version", "River [flood]", ""],
        "Author Keywords": ["prediction", "", "", "Water", ""],
    })


@pytest.mark.parametrize("value", ["https://doi.org/10.1000/ABC", "http://doi.org/10.1000/ABC", "https://dx.doi.org/10.1000/ABC", "doi:10.1000/ABC", "DOI: 10.1000/ABC", "10.1000/ABC"])
def test_normalize_doi(value):
    assert normalize_doi(value) == "10.1000/abc"


@pytest.mark.parametrize("value", [None, float("nan"), pd.NA, "nan", "None", "", "  "])
def test_missing_values(value):
    assert normalize_text(value) == normalize_doi(value) == ""


@pytest.mark.parametrize("dtype", ["object", "string[pyarrow]"])
def test_unicode_scientific_whitespace_normalizes_consistently(dtype):
    value = "\u00a0AUC\u2009improved\nby 20.67\u202f%\u3000"
    assert clean_series(pd.Series([value], dtype=dtype)).iloc[0] == normalize_text(value) == "AUC improved by 20.67 %"


@pytest.mark.parametrize("value,expected", [("2025-03-20", "2025"), ("April 2025", "2025"), ("20/04/2025", "2025"), (2025.0, "2025"), ("2025–2026", ""), ("2025 / 2025", "2025"), ("120251", ""), (None, "")])
def test_years(value, expected):
    assert extract_year(value) == expected


@pytest.mark.parametrize("separator", [",", ";", "\t"])
@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1252", "latin-1"])
def test_csv_encodings_and_separators(separator, encoding):
    source = pd.DataFrame({"Item Title": ["Predicción de lluvias"], "Authors": ["Pérez, Ana"], "Abstract": ["Texto; con, puntuación\ny saltos"]})
    content = source.to_csv(index=False, sep=separator).encode(encoding)
    parsed = read_uploaded_file(content, "export.csv")
    result = process_dataframe(parsed)
    assert result.clean.loc[0, "Titulo"] == "Predicción de lluvias"
    assert result.clean.loc[0, "Autores"] == "Pérez, Ana"
    assert result.clean.loc[0, "Resumen"] == "Texto; con, puntuación y saltos"
    assert parsed.attrs["separator"] == separator


def test_windows_specific_characters():
    parsed = read_uploaded_file('Title;Authors\n“Climate”;Muñoz'.encode("cp1252"), "export.csv")
    assert parsed.iloc[0, 0] == "“Climate”"


def test_latin1_fallback():
    parsed = read_uploaded_file(b"Title\nStudy \x81", "export.csv")
    assert parsed.attrs["encoding"] == "latin-1"


def test_excel_with_empty_first_sheet(bibliography):
    content = BytesIO()
    with pd.ExcelWriter(content, engine="openpyxl") as writer:
        pd.DataFrame().to_excel(writer, sheet_name="Vacía", index=False)
        bibliography.to_excel(writer, sheet_name="Bibliografía", index=False)
    parsed = read_uploaded_file(content.getvalue(), "export.XLSX")
    assert len(parsed) == 5
    assert parsed.attrs["sheet"] == "Bibliografía"
    assert len(process_dataframe(parsed).clean) == 3


def test_duplicates_and_prisma_schema(bibliography):
    result = process_dataframe(bibliography)
    assert result.records.Duplicado.tolist() == [False, True, True, False, False]
    assert result.duplicates.Motivo_duplicado.tolist() == ["DOI y título duplicados", "DOI duplicado"]
    assert result.clean.columns.tolist() == ["ID", *FIELDS, *SCREENING_FIELDS]
    assert result.clean.ID.tolist() == [1, 2, 3]
    assert result.clean.loc[0, "Titulo"] == "Deep Learning for Debris Flow"
    assert result.clean.loc[0, "Autores"] == "Ana Pérez"
    assert result.clean[SCREENING_FIELDS].eq("").all().all()
    assert list(result.metrics.values()) == [5, 2, 3, 4, 4]
    assert bibliography.loc[0, "DOI"] == "https://doi.org/10.1000/ABC"


def test_title_only_duplicates_and_absent_metadata():
    result = process_dataframe(pd.DataFrame({"Title": ["Deep Learning for Debris Flow", "deep learning for debris-flow", "Other"]}))
    assert result.records.Duplicado.tolist() == [False, True, False]
    assert result.duplicates.loc[0, "Motivo_duplicado"] == "Título duplicado"
    assert result.clean[["DOI", "Resumen"]].eq("").all().all()
    assert len(result.warnings) == 2


def test_empty_keys_are_never_duplicates():
    result = process_dataframe(pd.DataFrame({"Authors": ["Ana", "Lee", "Pat"]}))
    assert len(result.clean) == 3
    assert result.metrics["Duplicados encontrados"] == 0


def test_blank_rows_and_real_nulls_are_removed():
    data = pd.DataFrame({"Title": [None, "Paper", "  ", pd.NA], "DOI": ["nan", None, "None", None]})
    assert len(process_dataframe(data).clean) == 1


def test_duplicate_headers_are_coalesced():
    content = b'Title,DOI,DOI,Abstract, Abstract ,Author Keywords,Index Keywords\nPaper,,10.1/x,,Summary,AI,water\n'
    result = process_dataframe(read_uploaded_file(content, "export.csv"))
    assert result.clean.loc[0, "DOI"] == "10.1/x"
    assert result.clean.loc[0, "Resumen"] == "Summary"
    assert result.clean.loc[0, "Palabras_clave"] == "AI; water"
    assert len(result.mapping["DOI"]) == 2


def test_unique_header_suffix_collision():
    columns = unique_headers(["DOI", "DOI", "DOI [columna 2]", "DOI"])
    assert len(set(columns)) == 4


@pytest.mark.parametrize("headers", [
    [" Item Title ", "Author(s)", "Publication Date", "DOI", "Abstract Note", "Key Words", "Publication Title", "Web Link", "Content Type", "Affiliations"],
    ["Título", "Autores", "AÑO", "Document DOI", "Resumen", "Palabras_clave", "Revista", "URL", "Tipo_documento", "Afiliación"],
    ["TI", "AU", "PY", "DI", "AB", "DE", "SO", "UR", "DT", "C1"],
])
def test_aliases(headers):
    mapping = detect_columns(pd.DataFrame(columns=headers))
    assert all(mapping[field] for field in FIELDS)


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_springer_export_columns_and_doi_duplicates(extension):
    source = pd.DataFrame({
        "Item Title": ["Rainfall prediction", "Another version", "Flood modelling"],
        "Publication Title": ["Natural Hazards"] * 3,
        "Book Series Title": ["Environmental Sciences"] * 3,
        "Journal   Volume": ["125"] * 3,
        "Journal Issue": ["4"] * 3,
        "Item DOI": ["https://doi.org/10.1000/ABC", "DOI: 10.1000/abc", "10.1000/def"],
        "Authors": ["Ana Pérez", "Ana Pérez", "Lee"],
        "Publication Year": ["2025", "2025", "2024"],
        "URL": ["https://example.org/paper-a", "https://example.org/version", "https://example.org/paper-b"],
        "Content Type": ["Article"] * 3,
    })
    if extension == "csv":
        content = source.to_csv(index=False).encode("utf-8-sig")
    else:
        buffer = BytesIO()
        source.to_excel(buffer, index=False)
        content = buffer.getvalue()
    result = process_dataframe(read_uploaded_file(content, f"springer.{extension}"))
    expected_mapping = {
        "Titulo": "Item Title", "Revista": "Publication Title", "DOI": "Item DOI",
        "Autores": "Authors", "Año": "Publication Year", "URL": "URL",
        "Tipo_documento": "Content Type",
    }
    for field, column in expected_mapping.items():
        assert result.mapping[field] == [column]
    assert result.clean.loc[0, "Titulo"] == "Rainfall prediction"
    assert result.clean.loc[0, "Revista"] == "Natural Hazards"
    assert result.clean.loc[0, "DOI"] == "10.1000/abc"
    assert result.clean.loc[0, "Tipo_documento"] == "Article"
    assert result.clean.loc[0, "Año"] == "2025"
    assert result.clean[["Resumen", "Palabras_clave", "Afiliacion"]].eq("").all().all()
    assert result.records.Duplicado.tolist() == [False, True, False]
    assert result.duplicates.Motivo_duplicado.tolist() == ["DOI duplicado"]
    assert result.metrics["Registros con DOI"] == 3
    assert len(result.clean) == 2
    assert not any("columna de DOI" in warning for warning in result.warnings)


def test_unknown_columns_can_be_mapped():
    source = pd.DataFrame({"nombre raro": ["Paper"], "texto largo": ["Summary"]})
    with pytest.raises(ColumnDetectionError):
        process_dataframe(source)
    result = process_dataframe(source, {"Titulo": "nombre raro", "Resumen": "texto largo"})
    assert result.clean.loc[0, "Titulo"] == "Paper"
    assert result.clean.loc[0, "Resumen"] == "Summary"


def test_filters_are_literal_case_insensitive_and_non_destructive(bibliography):
    result = process_dataframe(bibliography)
    original = result.clean.copy(deep=True)
    assert len(filter_records(result.clean, query="NEURAL")) == 1
    assert len(filter_records(result.clean, query="water")) == 1
    assert len(filter_records(result.clean, query="[flood]")) == 1
    assert len(filter_records(result.clean, doi="Sin DOI", abstract="Con abstract")) == 1
    assert len(filter_records(result.clean, year_range=(2025, 2025))) == 2
    assert len(filter_records(result.clean, year_range=(2025, 2025), include_unknown_year=False)) == 1
    assert len(filter_records(result.clean, abstract="Sin abstract")) == 1
    assert filter_records(result.clean, query="absent").empty
    pd.testing.assert_frame_equal(result.clean, original)


def test_exports_roundtrip_and_summary(bibliography):
    result = process_dataframe(bibliography)
    csv_bytes = export_csv(result.clean)
    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    loaded = pd.read_csv(BytesIO(csv_bytes), encoding="utf-8-sig", keep_default_na=False)
    assert loaded.Titulo.tolist() == result.clean.Titulo.tolist()
    workbook = load_workbook(BytesIO(export_excel(result.clean, result.metrics)))
    assert workbook.sheetnames == ["Papers", "Resumen_PRISMA"]
    assert workbook["Papers"].max_row == 4
    assert workbook["Papers"]["B2"].value == result.clean.loc[0, "Titulo"]
    assert workbook["Resumen_PRISMA"]["B2"].value == 5
    duplicate_workbook = load_workbook(BytesIO(export_excel(result.duplicates, result.metrics, "Duplicados")))
    rows = list(duplicate_workbook["Duplicados"].values)
    assert rows[1][-1] == "DOI y título duplicados"


def test_exports_do_not_execute_formulas():
    data = process_dataframe(pd.DataFrame({"Title": ["=1+1", "+1+2", "@SUM(A1)"]})).clean
    workbook = load_workbook(BytesIO(export_excel(data)))
    assert workbook["Papers"]["B2"].data_type == "s"
    assert workbook["Papers"]["B2"].value == "=1+1"
    assert "'=1+1" in export_csv(data).decode("utf-8-sig")


def test_excel_does_not_silently_truncate_long_abstract():
    data = process_dataframe(pd.DataFrame({"Title": ["Paper"], "Abstract": ["x" * 32768]})).clean
    with pytest.raises(FileProcessingError, match="32.767"):
        export_excel(data)
    assert "x" * 32768 in export_csv(data).decode("utf-8-sig")


@pytest.mark.parametrize("content,name", [
    (b"", "empty.csv"), (b"Title,DOI\n", "headers.csv"),
    (b'Title,DOI\n"broken', "bad.csv"), (b"Title,DOI\nPaper,a,b", "ragged.csv"),
    (b"Title,DOI\nPaper,a\nOther", "short.csv"),
    (b"not an excel", "corrupt.xlsx"), (b"\x00\x01binary", "binary.csv"),
    (b"Title\nPaper", "unsupported.txt"), (b"Title,DOI\nNone,nan", "null.csv"),
])
def test_bad_input_has_actionable_error(content, name):
    with pytest.raises(FileProcessingError):
        read_uploaded_file(content, name)


def test_excel_without_data():
    content = BytesIO()
    pd.DataFrame(columns=["Title"]).to_excel(content, index=False)
    with pytest.raises(FileProcessingError, match="no contiene datos"):
        read_uploaded_file(content.getvalue(), "empty.xlsx")


def test_excel_separator_hint():
    parsed = read_uploaded_file(b"sep=;\nTitle;Year\nPaper;2025", "export.csv")
    assert process_dataframe(parsed).clean.loc[0, "Año"] == "2025"


def test_ten_thousand_records():
    source = pd.DataFrame({"Title": [f"Paper {index}" for index in range(9000)] + [f"Paper {index}" for index in range(1000)], "Year": "2025", "Abstract": "A scientific abstract."})
    result = process_dataframe(source)
    assert result.metrics["Registros identificados"] == 10000
    assert result.metrics["Registros únicos"] == 9000
    assert result.metrics["Duplicados encontrados"] == 1000
